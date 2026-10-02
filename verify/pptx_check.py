"""Open a .pptx with python-pptx (an independent OOXML reader) and report its structure as JSON."""
import json, sys
from pptx import Presentation
from pptx.util import Emu

p = Presentation(sys.argv[1])
out = {"slide_w_in": round(Emu(p.slide_width).inches, 3), "slide_h_in": round(Emu(p.slide_height).inches, 3), "slides": []}
for s in p.slides:
    info = {"pictures": 0, "text_boxes": 0, "shapes": 0, "texts": [], "fonts": set(), "cropped": 0}
    for sh in s.shapes:
        if sh.shape_type == 13:  # PICTURE
            info["pictures"] += 1
            if any(getattr(sh, f"crop_{k}") for k in ("left", "right", "top", "bottom")): info["cropped"] += 1
        elif sh.has_text_frame and sh.text_frame.text.strip():
            info["text_boxes"] += 1
            info["texts"].append(sh.text_frame.text)
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    if r.font.name: info["fonts"].add(r.font.name)
        else:
            info["shapes"] += 1
    info["fonts"] = sorted(info["fonts"])
    out["slides"].append(info)
print(json.dumps(out))
